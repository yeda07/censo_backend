from datetime import date
from io import BytesIO

from django.contrib.auth import get_user_model
from django.test import TestCase
from openpyxl import load_workbook
from rest_framework.test import APIClient

from censos.models import Censo
from familias.models import Familia, MetaCensoFamiliar
from personas.models import Persona


class FamiliaWorkflowTests(TestCase):
    def setUp(self):
        self.client = APIClient()
        self.client.force_authenticate(user=get_user_model().objects.create_user(username='tester', password='pass'))

    def persona(self, familia, documento):
        return Persona.objects.create(
            familida_id=familia, nombres='Nombre', apellidos=str(documento),
            tipo_documento='CC', numero_documento=documento,
            exp_documento=date(2020, 1, 1), fecha_nacimiento=date(2000, 1, 1),
            parentesco='HI', sexo='F', estado_civil='Soltero', profesion='Ninguna',
            escolaridad='PR', integrantes=2, direccion='Calle 1', telefono='123', usuario='tester',
        )

    def test_meta_y_progreso_distinct_por_vigencia(self):
        response = self.client.post('/familia/', {
            'numero_familia': 101, 'nombre_flia': 'Familia Uno',
            'vigencia': '2026', 'integrantes_previstos': 2,
        })
        self.assertEqual(response.status_code, 201)
        familia = Familia.objects.get(pk=response.data['id'])
        self.assertEqual(MetaCensoFamiliar.objects.get(familia=familia).integrantes_previstos, 2)
        url = f'/familia/{familia.id}/progreso/?vigencia=2026'
        self.assertEqual(self.client.get(url).data['censados'], 0)
        uno = self.persona(familia, 1001)
        Censo.objects.create(persona=uno, vigencia='2026')
        Censo.objects.create(persona=uno, vigencia='2026')
        self.assertEqual(self.client.get(url).data['censados'], 1)
        dos = self.persona(familia, 1002)
        Censo.objects.create(persona=dos, vigencia='2025')
        self.assertEqual(self.client.get(url).data['censados'], 1)
        Censo.objects.create(persona=dos, vigencia='2026')
        self.assertTrue(self.client.get(url).data['completo'])
        self.assertEqual(self.client.get(url).data['restantes'], 0)

    def test_creacion_antigua_y_reporte(self):
        empty = self.client.post('/familia/', {'numero_familia': 201, 'nombre_flia': '=Familia vacia'})
        self.assertEqual(empty.status_code, 201)
        familia = Familia.objects.create(numero_familia=202, nombre_flia='Familia Dos')
        uno = self.persona(familia, 2001)
        self.persona(familia, 2002)
        Censo.objects.create(persona=uno, vigencia='2026')

        excel = self.client.get('/familia/reporte/?formato=xlsx')
        self.assertEqual(excel.status_code, 200)
        book = load_workbook(BytesIO(excel.content))
        self.assertEqual(book['Resumen']['B2'].value, 2)
        self.assertEqual(book['Resumen']['B3'].value, 2)
        self.assertEqual(book['Resumen']['B6'].data_type, 's')
        self.assertEqual(book['Resumen']['D7'].value, 2)
        self.assertEqual(book['Miembros'].max_row, 3)

        filtered = self.client.get('/familia/reporte/?formato=xlsx&vigencia=2026')
        self.assertEqual(load_workbook(BytesIO(filtered.content))['Resumen']['B3'].value, 1)
        pdf = self.client.get('/familia/reporte/?formato=pdf')
        self.assertEqual(pdf.status_code, 200)
        self.assertTrue(pdf.content.startswith(b'%PDF'))

    def test_reporte_requiere_autenticacion(self):
        self.client.force_authenticate(user=None)
        self.assertIn(self.client.get('/familia/reporte/?formato=pdf').status_code, (401, 403))
