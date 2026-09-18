from io import BytesIO
from xml.sax.saxutils import escape

from django.db.models import Subquery
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from reportlab.lib import colors
from reportlab.lib.enums import TA_CENTER
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import mm
from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

from censos.models import Censo
from familias.models import Familia, MetaCensoFamiliar
from personas.models import Persona


def datos_familias(vigencia=None):
    familias = list(Familia.objects.order_by('numero_familia', 'id'))
    ids = [familia.id for familia in familias]
    personas = Persona.objects.filter(familida_id_id__in=ids).order_by('apellidos', 'nombres', 'id')
    if vigencia:
        censados = Censo.objects.filter(vigencia=vigencia).exclude(persona_id=None).values('persona_id')
        personas = personas.filter(id__in=Subquery(censados))

    por_familia = {familia_id: [] for familia_id in ids}
    for persona in personas:
        por_familia[persona.familida_id_id].append(persona)

    metas = {}
    if vigencia:
        metas = {
            meta.familia_id: meta.integrantes_previstos
            for meta in MetaCensoFamiliar.objects.filter(familia_id__in=ids, vigencia=vigencia)
        }

    filas = [
        {
            'familia': familia,
            'miembros': por_familia[familia.id],
            'previstos': metas.get(familia.id),
        }
        for familia in familias
    ]
    return {
        'vigencia': vigencia,
        'familias': filas,
        'total_familias': len(filas),
        'total_miembros': sum(len(fila['miembros']) for fila in filas),
    }


def _texto_celda(hoja, fila, columna, valor):
    celda = hoja.cell(fila, columna, '' if valor is None else str(valor))
    celda.data_type = 's'
    return celda


def generar_excel(datos):
    libro = Workbook()
    resumen = libro.active
    resumen.title = 'Resumen'
    detalle = libro.create_sheet('Miembros')

    resumen.append(['Vigencia', datos['vigencia'] or 'Todas'])
    resumen.append(['Total familias', datos['total_familias']])
    resumen.append(['Total miembros', datos['total_miembros']])
    resumen.append([])
    resumen.append(['Número familia', 'Familia', 'Meta', 'Miembros'])
    detalle.append(['Número familia', 'Familia', 'Nombres', 'Apellidos', 'Tipo documento', 'Documento', 'Parentesco'])

    for indice, fila in enumerate(datos['familias'], start=6):
        familia = fila['familia']
        resumen.cell(indice, 1, familia.numero_familia)
        _texto_celda(resumen, indice, 2, familia.nombre_flia)
        resumen.cell(indice, 3, fila['previstos'])
        resumen.cell(indice, 4, len(fila['miembros']))
        for persona in fila['miembros']:
            linea = detalle.max_row + 1
            detalle.cell(linea, 1, familia.numero_familia)
            for columna, valor in enumerate([
                familia.nombre_flia, persona.nombres, persona.apellidos,
                persona.tipo_documento, persona.numero_documento, persona.get_parentesco_display(),
            ], start=2):
                _texto_celda(detalle, linea, columna, valor)
            detalle.cell(linea, 6).number_format = '@'

    for hoja, cabecera, anchos in (
        (resumen, 5, [19, 32, 12, 14]),
        (detalle, 1, [19, 32, 24, 24, 19, 22, 22]),
    ):
        for columna, ancho in enumerate(anchos, start=1):
            hoja.column_dimensions[hoja.cell(cabecera, columna).column_letter].width = ancho
        for celda in hoja[cabecera]:
            celda.fill = PatternFill('solid', fgColor='1877F2')
            celda.font = Font(color='FFFFFF', bold=True)
            celda.alignment = Alignment(wrap_text=True)
        hoja.freeze_panes = f'A{cabecera + 1}'
        hoja.auto_filter.ref = f'A{cabecera}:{hoja.cell(max(hoja.max_row, cabecera), len(anchos)).coordinate}'

    salida = BytesIO()
    libro.save(salida)
    return salida.getvalue()


def generar_pdf(datos):
    salida = BytesIO()
    documento = SimpleDocTemplate(
        salida, pagesize=A4, leftMargin=18 * mm, rightMargin=18 * mm,
        topMargin=18 * mm, bottomMargin=18 * mm,
    )
    estilos = getSampleStyleSheet()
    titulo = ParagraphStyle('TituloReporte', parent=estilos['Heading1'], textColor=colors.HexColor('#1877F2'))
    subtitulo = ParagraphStyle('FamiliaReporte', parent=estilos['Heading3'], spaceBefore=12, keepWithNext=True)
    celda = ParagraphStyle('CeldaReporte', parent=estilos['Normal'], fontSize=8, leading=10)
    centrado = ParagraphStyle('CentroReporte', parent=celda, alignment=TA_CENTER)

    contenido = [
        Paragraph('Reporte de familias y miembros', titulo),
        Paragraph(f"Vigencia: {escape(datos['vigencia'] or 'Todas')}", estilos['Normal']),
        Paragraph(
            f"Familias: {datos['total_familias']}  |  Miembros: {datos['total_miembros']}",
            estilos['Normal'],
        ),
        Spacer(1, 8 * mm),
    ]

    for fila in datos['familias']:
        familia = fila['familia']
        miembros = fila['miembros']
        nombre = escape(familia.nombre_flia or '')
        titulo_familia = f"Familia {familia.numero_familia or '-'}: {nombre} ({len(miembros)} miembros)"
        if fila['previstos'] is not None:
            titulo_familia += f" - Meta: {fila['previstos']}"
        contenido.append(Paragraph(titulo_familia, subtitulo))
        if not miembros:
            contenido.append(Paragraph('Sin miembros para este reporte.', estilos['Normal']))
            continue
        tabla = [[
            Paragraph('Nombre', centrado), Paragraph('Documento', centrado),
            Paragraph('Parentesco', centrado),
        ]]
        for persona in miembros:
            tabla.append([
                Paragraph(escape(f'{persona.nombres} {persona.apellidos}'), celda),
                Paragraph(escape(f'{persona.tipo_documento or ""} {persona.numero_documento}'), celda),
                Paragraph(escape(persona.get_parentesco_display() or ''), celda),
            ])
        tabla_pdf = Table(tabla, colWidths=[71 * mm, 51 * mm, 52 * mm], repeatRows=1, hAlign='LEFT')
        tabla_pdf.setStyle(TableStyle([
            ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('#E8F1FE')),
            ('VALIGN', (0, 0), (-1, -1), 'TOP'),
            ('BOTTOMPADDING', (0, 0), (-1, -1), 6),
            ('TOPPADDING', (0, 0), (-1, -1), 6),
            ('LINEBELOW', (0, -1), (-1, -1), 0.4, colors.HexColor('#DFE3E8')),
        ]))
        contenido.append(tabla_pdf)

    def numero_pagina(canvas, doc):
        canvas.saveState()
        canvas.setFont('Helvetica', 8)
        canvas.drawRightString(A4[0] - 18 * mm, 10 * mm, f'Página {doc.page}')
        canvas.restoreState()

    documento.build(contenido, onFirstPage=numero_pagina, onLaterPages=numero_pagina)
    return salida.getvalue()
