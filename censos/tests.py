from django.test import TestCase

# Create your tests here.
from io import BytesIO, StringIO

from openpyxl import load_workbook

from django.core.management import call_command
from django.test import TestCase, override_settings
from django.utils import timezone
from rest_framework.test import APIClient

from actividades.models import Actividad
from censo_actividades.models import CensoActividad
from censos.models import Censo
from censos.reportes import EXCEL_FIELDS
from familias.models import Familia, MetaCensoFamiliar
from multas.models import Multa
from personas.models import Persona
from users.models import User


@override_settings(DEBUG=True)
class DatosPruebaTests(TestCase):
    def test_carga_repetible_y_estados(self):
        output = StringIO()
        call_command('cargar_datos_prueba', stdout=output)
        call_command('cargar_datos_prueba', stdout=output)

        self.assertEqual(Familia.objects.count(), 2)
        self.assertEqual(Persona.objects.count(), 3)
        self.assertEqual(Censo.objects.count(), 3)
        self.assertEqual(Actividad.objects.count(), 3)
        self.assertEqual(CensoActividad.objects.count(), 3)
        self.assertEqual(Multa.objects.count(), 1)

        estados = {
            item.censo.persona.nombres: item.estado
            for item in CensoActividad.objects.select_related('censo__persona')
        }
        self.assertEqual(estados, {
            'Ana DEMO': CensoActividad.TERMINADA,
            'Bruno DEMO': CensoActividad.PENDIENTE,
            'Carla DEMO': CensoActividad.NO_REALIZADA,
        })

    def test_personas_se_filtran_por_familia(self):
        call_command('cargar_datos_prueba', stdout=StringIO())
        user = User.objects.create_user(
            username='operador', email='operador@example.test', password='TestPassword123!'
        )
        client = APIClient()
        client.force_authenticate(user=user)

        familia_a = Familia.objects.get(numero_familia=90001)
        familia_b = Familia.objects.get(numero_familia=90002)
        self.assertEqual(len(client.get(f'/persona/?familia_id={familia_a.id}').data), 2)
        self.assertEqual(len(client.get(f'/persona/?familia_id={familia_b.id}').data), 1)

    def test_reporte_general_censo_incluye_todos_los_campos_y_filtra_vigencia(self):
        call_command('cargar_datos_prueba', stdout=StringIO())
        user = User.objects.create_user(
            username='reportes', email='reportes@example.test', password='TestPassword123!'
        )
        client = APIClient()
        self.assertEqual(client.get('/censo/reporte/?formato=xlsx').status_code, 401)
        client.force_authenticate(user=user)
        vigencia = Censo.objects.first().vigencia
        esperado = Censo.objects.filter(vigencia=vigencia).count()

        respuesta = client.get(f'/censo/reporte/?formato=xlsx&vigencia={vigencia}')
        self.assertEqual(respuesta.status_code, 200)
        self.assertIn('attachment;', respuesta['Content-Disposition'])
        libro = load_workbook(BytesIO(respuesta.content), read_only=True)
        hoja = libro['Censo poblacional']
        filas = list(hoja.values)
        self.assertEqual(filas[0], tuple(label for label, _ in EXCEL_FIELDS))
        self.assertEqual(len(filas) - 1, esperado)
        self.assertTrue(all(fila[0] == vigencia for fila in filas[1:]))
        self.assertIn('Discapacidad', filas[0])
        for columna in ('ID censo', 'ID persona', 'ID familia', 'Usuario'):
            self.assertNotIn(columna, filas[0])

        pdf = client.get(f'/censo/reporte/?formato=pdf&vigencia={vigencia}')
        self.assertEqual(pdf.status_code, 200)
        self.assertTrue(pdf.content.startswith(b'%PDF-'))
        self.assertGreater(len(pdf.content), 1500)

        vacio = client.get('/censo/reporte/?formato=xlsx&vigencia=sin-registros')
        self.assertEqual(vacio.status_code, 200)
        self.assertEqual(len(list(load_workbook(BytesIO(vacio.content), read_only=True)['Censo poblacional'].values)), 1)
        self.assertEqual(client.get('/censo/reporte/?formato=csv').status_code, 400)

    def test_otra_vigencia_reutiliza_persona(self):
        call_command('cargar_datos_prueba', stdout=StringIO())
        user = User.objects.create_user(
            username='operador', email='operador@example.test', password='TestPassword123!'
        )
        client = APIClient()
        client.force_authenticate(user=user)
        persona = Persona.objects.get(nombres='Ana DEMO')

        response = client.post(
            '/censo/', {'persona_id': persona.id, 'vigencia': '2027'}, format='json'
        )
        self.assertEqual(response.status_code, 201, response.data)
        self.assertEqual(response.data['persona']['id'], persona.id)
        self.assertEqual(Persona.objects.count(), 3)
        self.assertEqual(
            client.post('/censo/', {'persona_id': persona.id, 'vigencia': '2027'}, format='json').status_code,
            400,
        )

    def test_crear_censo_con_campos_de_persona_existente(self):
        call_command('cargar_datos_prueba', stdout=StringIO())
        user = User.objects.create_user(
            username='operador', email='operador@example.test', password='TestPassword123!'
        )
        client = APIClient()
        client.force_authenticate(user=user)
        persona = Persona.objects.get(nombres='Ana DEMO')
        datos = {
            'nombres': persona.nombres,
            'apellidos': 'Actualizado DEMO',
            'tipo_documento': persona.tipo_documento,
            'numero_documento': persona.numero_documento,
            'exp_documento': str(persona.exp_documento),
            'fecha_nacimiento': str(persona.fecha_nacimiento),
            'parentesco': persona.parentesco,
            'sexo': persona.sexo,
            'estado_civil': persona.estado_civil,
            'profesion': persona.profesion,
            'escolaridad': persona.escolaridad,
            'integrantes': persona.integrantes,
            'direccion': persona.direccion,
            'telefono': persona.telefono,
            'usuario': persona.usuario,
            'familida_id': persona.familida_id_id,
        }
        response = client.post('/censo/', {
            'persona_id': persona.id, 'vigencia': '2027', 'persona_data': datos,
        }, format='json')
        self.assertEqual(response.status_code, 201, response.data)
        persona.refresh_from_db()
        self.assertEqual(persona.apellidos, 'Actualizado DEMO')
        self.assertEqual(response.data['persona']['id'], persona.id)
        self.assertEqual(Persona.objects.count(), 3)

        datos['apellidos'] = 'No debe guardarse'
        duplicado = client.post('/censo/', {
            'persona_id': persona.id, 'vigencia': '2027', 'persona_data': datos,
        }, format='json')
        self.assertEqual(duplicado.status_code, 400)
        persona.refresh_from_db()
        self.assertEqual(persona.apellidos, 'Actualizado DEMO')

        actualizado = client.patch(f'/censo/{response.data["id"]}/', {
            'vigencia': '2028', 'persona_data': {**datos, 'apellidos': 'Editado DEMO'},
        }, format='json')
        self.assertEqual(actualizado.status_code, 200, actualizado.data)
        persona.refresh_from_db()
        self.assertEqual(persona.apellidos, 'Editado DEMO')
        self.assertEqual(actualizado.data['vigencia'], '2028')

        nueva_persona = {
            **datos, 'nombres': 'Nueva DEMO', 'apellidos': 'Persona DEMO',
            'numero_documento': 999001, 'tipo_documento': 'CE', 'discapacidad': 'SI',
        }
        nuevo_censo = client.post('/censo/', {
            'vigencia': '2027', 'resguardo_ind': 1, 'comunidad_ind': 1,
            'persona': nueva_persona,
        }, format='json')
        self.assertEqual(nuevo_censo.status_code, 201, nuevo_censo.data)
        self.assertEqual(Persona.objects.count(), 4)
        self.assertEqual(nuevo_censo.data['persona']['nombres'], 'Nueva DEMO')
        self.assertEqual(nuevo_censo.data['persona']['tipo_documento'], 'CE')
        self.assertEqual(nuevo_censo.data['persona']['discapacidad'], 'SI')
        self.assertEqual(nuevo_censo.data['persona']['familida_id'], nueva_persona['familida_id'])
        self.assertTrue(Censo.objects.filter(pk=nuevo_censo.data['id'], persona_id=nuevo_censo.data['persona']['id']).exists())


@override_settings(DEBUG=True)
class DatosPruebaAmpliadosTests(TestCase):
    def test_volumen_estados_y_carga_repetible(self):
        real = Familia.objects.create(numero_familia=22, nombre_flia='Familia existente')
        call_command('cargar_datos_prueba_ampliados', stdout=StringIO())
        conteos = (
            Familia.objects.count(), Persona.objects.count(), Censo.objects.count(),
            Actividad.objects.count(), CensoActividad.objects.count(), Multa.objects.count(),
        )
        self.assertEqual(conteos[:2], (35, 105))
        self.assertGreater(conteos[2], 110)
        self.assertEqual(conteos[3:], (63, 93, 31))
        self.assertEqual(Persona.objects.filter(familida_id__numero_familia=91001).count(), 12)
        self.assertEqual(Persona.objects.filter(familida_id__numero_familia=91032).count(), 0)
        self.assertEqual(MetaCensoFamiliar.objects.count(), 32)
        self.assertGreater(Persona.objects.filter(discapacidad='SI').count(), 0)
        self.assertEqual(CensoActividad.objects.filter(estado=CensoActividad.TERMINADA).count(), 31)
        self.assertEqual(CensoActividad.objects.filter(estado=CensoActividad.PENDIENTE).count(), 31)
        self.assertEqual(CensoActividad.objects.filter(estado=CensoActividad.NO_REALIZADA).count(), 31)

        call_command('cargar_datos_prueba_ampliados', stdout=StringIO())
        self.assertEqual((
            Familia.objects.count(), Persona.objects.count(), Censo.objects.count(),
            Actividad.objects.count(), CensoActividad.objects.count(), Multa.objects.count(),
        ), conteos)
        self.assertEqual(Familia.objects.get(pk=real.pk).nombre_flia, 'Familia existente')

    def test_progreso_reportes_y_paz_y_salvo(self):
        call_command('cargar_datos_prueba_ampliados', stdout=StringIO())
        user = User.objects.create_user(username='operador', email='operador@example.test', password='TestPassword123!')
        client = APIClient()
        client.force_authenticate(user=user)
        vigencia = f'DEMO-{timezone.localdate().year}'

        parcial = Familia.objects.get(numero_familia=91002)
        progreso = client.get(f'/familia/{parcial.id}/progreso/?vigencia={vigencia}')
        self.assertEqual(progreso.data['censados'], 3)
        self.assertEqual(progreso.data['integrantes_previstos'], 5)
        self.assertFalse(progreso.data['completo'])
        vacia = Familia.objects.get(numero_familia=91032)
        self.assertEqual(client.get(f'/familia/{vacia.id}/progreso/?vigencia={vigencia}').data['censados'], 0)

        reporte = client.get('/familia/reporte/?formato=pdf')
        self.assertEqual(reporte.status_code, 200)
        self.assertTrue(reporte.content.startswith(b'%PDF'))
        excel = client.get('/familia/reporte/?formato=xlsx')
        self.assertEqual(excel.status_code, 200)
        self.assertGreater(len(excel.content), 1000)

        ana = Censo.objects.get(persona__nombres='Ana DEMO', vigencia=vigencia)
        carla = Censo.objects.get(persona__nombres='Carla DEMO', vigencia=vigencia)
        self.assertTrue(client.get(f'/censo/{ana.id}/paz-y-salvo/').data['habilitado'])
        self.assertFalse(client.get(f'/censo/{carla.id}/paz-y-salvo/').data['habilitado'])
