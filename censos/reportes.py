from io import BytesIO
from xml.sax.saxutils import escape

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import mm
from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

from censos.models import Censo


FIELDS = [
    ('ID censo', 'id'),
    ('Vigencia', 'vigencia'),
    ('Resguardo indígena', 'resguardo_ind'),
    ('Comunidad indígena', 'comunidad_ind'),
    ('ID persona', 'persona.id'),
    ('ID familia', 'persona.familida_id_id'),
    ('Número familia', 'persona.familida_id.numero_familia'),
    ('Familia', 'persona.familida_id.nombre_flia'),
    ('Nombres', 'persona.nombres'),
    ('Apellidos', 'persona.apellidos'),
    ('Tipo documento', 'persona.tipo_documento'),
    ('Número documento', 'persona.numero_documento'),
    ('Expedición documento', 'persona.exp_documento'),
    ('Fecha nacimiento', 'persona.fecha_nacimiento'),
    ('Parentesco', 'persona.parentesco'),
    ('Sexo', 'persona.sexo'),
    ('Estado civil', 'persona.estado_civil'),
    ('Profesión', 'persona.profesion'),
    ('Escolaridad', 'persona.escolaridad'),
    ('Discapacidad', 'persona.discapacidad'),
    ('Integrantes familia', 'persona.integrantes'),
    ('Dirección', 'persona.direccion'),
    ('Teléfono', 'persona.telefono'),
    ('Usuario', 'persona.usuario'),
]
EXCEL_FIELDS = [
    field for field in FIELDS
    if field[1] not in {'id', 'persona.id', 'persona.familida_id_id', 'persona.usuario'}
]


def datos_censos(vigencia=None):
    censos = Censo.objects.select_related('persona__familida_id').order_by('vigencia', 'id')
    if vigencia:
        censos = censos.filter(vigencia=vigencia)
    return censos


def _valor(censo, path):
    objeto = censo
    for atributo in path.split('.'):
        if objeto is None:
            return ''
        objeto = getattr(objeto, atributo, None)
    if objeto is None:
        return ''
    if path in ('persona.tipo_documento', 'persona.parentesco', 'persona.escolaridad', 'persona.discapacidad'):
        display = getattr(censo.persona, f'get_{path.split(".")[-1]}_display')()
        return f'{display} ({objeto})' if display != objeto else str(objeto)
    if hasattr(objeto, 'isoformat'):
        return objeto.isoformat()
    return objeto


def generar_excel(censos, vigencia=None):
    libro = Workbook()
    resumen = libro.active
    resumen.title = 'Resumen'
    detalle = libro.create_sheet('Censo poblacional')
    registros = list(censos)
    resumen.append(['Vigencia', vigencia or 'Todas'])
    resumen.append(['Total registros', len(registros)])
    resumen.append(['Personas únicas', len({censo.persona_id for censo in registros if censo.persona_id})])
    detalle.append([label for label, _ in EXCEL_FIELDS])
    for censo in registros:
        detalle.append([_valor(censo, path) for _, path in EXCEL_FIELDS])
        for celda in detalle[detalle.max_row]:
            if isinstance(celda.value, str):
                celda.data_type = 's'

    for celda in detalle[1]:
        celda.fill = PatternFill('solid', fgColor='1877F2')
        celda.font = Font(color='FFFFFF', bold=True)
        celda.alignment = Alignment(wrap_text=True)
        detalle.column_dimensions[celda.column_letter].width = 25
    detalle.freeze_panes = 'A2'
    detalle.auto_filter.ref = detalle.dimensions
    resumen.column_dimensions['A'].width = 22
    resumen.column_dimensions['B'].width = 20
    salida = BytesIO()
    libro.save(salida)
    return salida.getvalue()


def generar_pdf(censos, vigencia=None):
    registros = list(censos)
    salida = BytesIO()
    documento = SimpleDocTemplate(
        salida, pagesize=A4, leftMargin=16 * mm, rightMargin=16 * mm,
        topMargin=16 * mm, bottomMargin=16 * mm,
    )
    estilos = getSampleStyleSheet()
    titulo = ParagraphStyle('TituloCenso', parent=estilos['Heading1'], textColor=colors.HexColor('#1877F2'))
    encabezado = ParagraphStyle('EncabezadoCenso', parent=estilos['Heading3'], spaceBefore=10, spaceAfter=5, keepWithNext=True)
    etiqueta = ParagraphStyle('EtiquetaCenso', parent=estilos['Normal'], fontSize=7, leading=9, textColor=colors.HexColor('#637381'))
    valor = ParagraphStyle('ValorCenso', parent=estilos['Normal'], fontSize=8, leading=10, wordWrap='CJK')
    contenido = [
        Paragraph('Censo poblacional general', titulo),
        Paragraph(f"Vigencia: {escape(vigencia or 'Todas')} | Registros: {len(registros)} | Personas únicas: {len({censo.persona_id for censo in registros if censo.persona_id})}", estilos['Normal']),
        Spacer(1, 4 * mm),
    ]
    if not registros:
        contenido.append(Paragraph('No hay censos para la vigencia seleccionada.', estilos['Normal']))

    ancho_etiqueta = 33 * mm
    ancho_valor = 56 * mm
    for censo in registros:
        persona = censo.persona
        nombre = f'{persona.nombres} {persona.apellidos}' if persona else 'Sin persona'
        contenido.append(Paragraph(
            escape(f'Censo {censo.id} | {nombre} | Vigencia {censo.vigencia}'), encabezado,
        ))
        pares = [(label, _valor(censo, path)) for label, path in FIELDS]
        filas = []
        for indice in range(0, len(pares), 2):
            fila = []
            for label, value in pares[indice:indice + 2]:
                fila.extend([Paragraph(escape(label), etiqueta), Paragraph(escape(str(value)), valor)])
            if len(fila) == 2:
                fila.extend(['', ''])
            filas.append(fila)
        tabla = Table(filas, colWidths=[ancho_etiqueta, ancho_valor, ancho_etiqueta, ancho_valor], hAlign='LEFT')
        tabla.setStyle(TableStyle([
            ('VALIGN', (0, 0), (-1, -1), 'TOP'),
            ('ROWBACKGROUNDS', (0, 0), (-1, -1), [colors.white, colors.HexColor('#F4F6F8')]),
            ('TOPPADDING', (0, 0), (-1, -1), 4),
            ('BOTTOMPADDING', (0, 0), (-1, -1), 4),
            ('LEFTPADDING', (0, 0), (-1, -1), 4),
            ('RIGHTPADDING', (0, 0), (-1, -1), 4),
        ]))
        contenido.append(tabla)

    def numero_pagina(canvas, doc):
        canvas.saveState()
        canvas.setFont('Helvetica', 8)
        canvas.drawRightString(A4[0] - 16 * mm, 9 * mm, f'Página {doc.page}')
        canvas.restoreState()

    documento.build(contenido, onFirstPage=numero_pagina, onLaterPages=numero_pagina)
    return salida.getvalue()
